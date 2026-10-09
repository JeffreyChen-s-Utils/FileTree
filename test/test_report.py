"""Reports preserve recorded coverage, bounded counts, escaped data and atomic cancellation."""

import threading

from openpyxl import load_workbook
import pytest

from je_file_tree.core.analysis import summarise
from je_file_tree.core.node import Node
from je_file_tree.core.report import Report, ReportCancelledError, ReportTable, prepare_report, write_html
from je_file_tree.core.scanner import ScanOptions, scan
from je_file_tree.gui.report_dialog import _labels
from je_file_tree.gui.report_export import write_xlsx


def test_recorded_scope_is_copied_without_original_file_reads(sample_tree, monkeypatch):
    root = scan(sample_tree).root
    summary = summarise(root)
    def forbidden(*_args, **_kwargs):
        pytest.fail("report preparation cannot read the filesystem")
    monkeypatch.setattr("os.stat", forbidden)
    monkeypatch.setattr("os.scandir", forbidden)
    report = prepare_report(root, summary, _labels())
    assert len(report.tables) == 6
    assert dict(report.tables[0].rows)["Logical bytes"] == 1000
    assert dict(report.tables[0].rows)["Notes"] == report.note
    assert report.tables[1].count == 3
    assert report.tables[2].count == 6 and report.tables[2].rows[0][1] == 500
    before = report.tables[2].rows[0]
    summary.largest[0].size = 0
    assert report.tables[2].rows[0] == before


def test_partial_and_excluded_scopes_are_not_complete(sample_tree):
    root = scan(sample_tree, options=ScanOptions(exclude=("code",))).root
    for partial in (False, True):
        report = prepare_report(root, summarise(root), _labels(), partial=partial)
        rows = dict(report.tables[0].rows)
        assert rows["Coverage"] == "Incomplete"
        assert rows["Skipped folders"] == 1
        assert rows["Logical bytes"] == 850


def test_deep_and_equal_sized_top_folders_are_bounded():
    root = Node("/recorded", True, children=[])
    parent = root
    for _ in range(1200):
        child = Node("深", True, size=1, children=[], parent=parent)
        parent.children.append(child)
        parent = child
    report = prepare_report(root, summarise(root), _labels())
    assert report.tables[1].count == 1200 and len(report.tables[1].rows) == 1000


def _unsafe_report():
    table = ReportTable('=unsafe/[table]', ("Path", "Bytes"),
                        (("<script>alert('x')</script>&資料", 123), ("=cmd", 2**60),
                         ("@formula\x01", 4), ("x" * 40000, 5)), 4)
    return Report("Report <unsafe>", "known & estimated", (table,))


def test_html_is_self_contained_escaped_and_has_no_executable_markup(tmp_path):
    target = tmp_path / "報告.html"
    assert write_html(_unsafe_report(), [("<chart>", b"\x89PNG\r\n\x1a\nexample")], str(target)) == 4
    text = target.read_text(encoding="utf-8")
    assert "<script>" not in text and "&lt;script&gt;" in text
    assert "&lt;unsafe&gt;" in text and "data:image/png;base64," in text
    assert 'default-src \'none\'' in text and "<chart>" not in text
    assert "https://" not in text


def test_excel_preserves_numbers_and_escapes_formulas_controls_and_cell_limits(tmp_path):
    target = tmp_path / "報告.xlsx"
    assert write_xlsx(_unsafe_report(), str(target)) == 4
    book = load_workbook(target)
    try:
        sheet = book.active
        assert sheet.freeze_panes == "A4" and sheet.auto_filter.ref == "A3:B7"
        assert sheet["B4"].value == 123 and sheet["B4"].data_type == "n"
        assert sheet["A5"].value == "'=cmd" and sheet["A5"].data_type == "s"
        assert sheet["B5"].value == str(2**60) and sheet["B5"].data_type == "s"
        assert sheet["A6"].value == "'@formula\\x01"
        assert len(sheet["A7"].value) == 32767
        assert not any(cell.data_type == "f" for row in sheet for cell in row)
    finally:
        book.close()


@pytest.mark.parametrize("kind", ("html", "xlsx"))
def test_cancel_and_failed_replace_preserve_previous_target(tmp_path, monkeypatch, kind):
    target = tmp_path / ("existing." + kind)
    target.write_bytes(b"previous")
    cancel = threading.Event()
    cancel.set()
    def write():
        if kind == "html":
            return write_html(_unsafe_report(), [], str(target), cancel=cancel)
        return write_xlsx(_unsafe_report(), str(target), cancel=cancel)
    with pytest.raises(ReportCancelledError):
        write()
    assert target.read_bytes() == b"previous"
    cancel.clear()
    def denied(*_args):
        raise PermissionError("fixture refuses replacement")
    monkeypatch.setattr("os.replace", denied)
    with pytest.raises(PermissionError):
        write()
    assert target.read_bytes() == b"previous"
    assert not list(tmp_path.glob(".file-tree-*"))


def test_cancel_before_preparation_and_invalid_png_are_errors(sample_tree, tmp_path):
    root = scan(sample_tree).root
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(ReportCancelledError):
        prepare_report(root, summarise(root), _labels(), cancel=cancel)
    with pytest.raises(ValueError, match="PNG"):
        write_html(_unsafe_report(), [("chart", b"not png")], str(tmp_path / "absent.html"))
    assert not (tmp_path / "absent.html").exists()
