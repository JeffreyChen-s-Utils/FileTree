"""Capture three independent chart views and write bounded report lists to Excel outside the core."""

from __future__ import annotations

import os
import re
import tempfile
import threading
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PySide6.QtCore import QByteArray, QBuffer, QIODevice, Qt
from PySide6.QtGui import QImage

from je_file_tree.core.export import spreadsheet_text
from je_file_tree.core.node import Node
from je_file_tree.core.report import Cell, Report, check_cancel
from je_file_tree.gui.charts import BARS, SUNBURST, TREEMAP, ChartStack
from je_file_tree.gui.i18n import tr

_XML_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_CELL_LIMIT = 32767
_EXACT_INTEGER_LIMIT = 10**15


def capture_report_charts(root: Node, source: ChartStack, unit: str, now: float) -> list[tuple[str, QImage]]:
    """GUI thread: grab whole-root Treemap, Bars and Sunburst without changing the user's views."""
    charts = ChartStack()
    charts.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen)
    charts.resize(1000, 600)
    charts.set_unit(unit)
    charts.set_colour_mode(source.treemap.colour_mode)
    charts.set_age_reference(now)
    charts.treemap.set_levels(source.treemap.levels)
    charts.set_view_root(root)
    charts.show()
    try:
        captured = []
        for mode in (TREEMAP, BARS, SUNBURST):
            charts.set_mode(mode)
            captured.append((tr("chart_" + mode), charts.currentWidget().grab().toImage()))
        return captured
    finally:
        charts.close()
        charts.deleteLater()


def encode_charts(charts: Sequence[tuple[str, QImage]], cancel: threading.Event) -> list[tuple[str, bytes]]:
    """Worker thread: encode owned chart images, observing cancellation between images."""
    encoded = []
    for title, image in charts:
        check_cancel(cancel)
        data = QByteArray()
        buffer = QBuffer(data)
        if not buffer.open(QIODevice.OpenModeFlag.WriteOnly) or not image.save(buffer, "PNG"):
            raise OSError("Report PNG encoding failed")
        buffer.close()
        encoded.append((title, bytes(data)))
    return encoded


def _excel_value(value: Cell) -> Cell:
    if isinstance(value, int) and abs(value) < _EXACT_INTEGER_LIMIT:
        return value
    text = _XML_CONTROL.sub(lambda match: f"\\x{ord(match[0]):02x}", str(value))
    return spreadsheet_text(text)[:_CELL_LIMIT]


def write_xlsx(report: Report, target: str, *, cancel: threading.Event | None = None) -> int:
    """Atomically write six sheets with byte/count numbers, escaped text and bounded cell lengths."""
    cancel = cancel if cancel is not None else threading.Event()
    handle, temporary = tempfile.mkstemp(prefix=".file-tree-", suffix=".xlsx",
                                         dir=os.path.dirname(os.path.abspath(target)))
    os.close(handle)
    workbook, count = Workbook(), 0
    workbook.remove(workbook.active)
    try:
        for number, table in enumerate(report.tables, 1):
            check_cancel(cancel)
            title = re.sub(r"[\\/*?:\[\]]", "_", table.title).strip("'")
            sheet = workbook.create_sheet((f"{number} " + title)[:31])
            sheet.freeze_panes = "A4"
            sheet.append([_excel_value(report.title)])
            sheet.append([_excel_value(f"{table.title} ({len(table.rows)} / {table.count})")])
            sheet.append([_excel_value(header) for header in table.header])
            for row in table.rows:
                check_cancel(cancel)
                sheet.append([_excel_value(value) for value in row])
                count += 1
            sheet.auto_filter.ref = f"A3:{sheet.cell(3, len(table.header)).column_letter}{sheet.max_row}"
            _style_sheet(sheet, cancel)
        check_cancel(cancel)
        workbook.save(temporary)
        check_cancel(cancel)
        os.replace(temporary, target)
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.unlink(temporary)
    return count


def _style_sheet(sheet, cancel: threading.Event) -> None:
    for cell in sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24466B")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(72, max(18, max(len(str(cell.value or ""))
                                                                   for cell in column[:60]) + 2))
    for row in sheet.iter_rows(min_row=4):
        check_cancel(cancel)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
